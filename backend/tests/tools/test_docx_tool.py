"""Tests for the DOCX document tool."""

import os
import tempfile
import pytest


@pytest.mark.asyncio
async def test_read_docx_file_not_found():
    """read_docx returns error for missing file."""
    from app.tools.docx_tool import read_docx
    result = await read_docx({"file_path": "/nonexistent/test.docx"}, {})
    assert result["error"] is not None
    assert "File not found" in result["error"]
    assert result["type"] == "document"
    assert result["tool"] == "read_docx"


@pytest.mark.asyncio
async def test_read_docx_empty_params():
    """read_docx returns error when no file_path given."""
    from app.tools.docx_tool import read_docx
    result = await read_docx({}, {})
    assert result["error"] is not None


@pytest.mark.asyncio
async def test_read_docx_valid_file():
    """read_docx extracts text from a valid .docx."""
    from app.tools.docx_tool import read_docx
    docx_path = _create_minimal_docx()
    if not docx_path:
        pytest.skip("python-docx not available")

    try:
        result = await read_docx({"file_path": docx_path}, {})
        assert result["error"] is None
        assert result["type"] == "document"
        assert result["tool"] == "read_docx"
        assert "Hello" in result["content"] or "test" in result["content"] or "Project" in result["content"]
        assert result["metadata"]["paragraph_count"] >= 1
    finally:
        os.unlink(docx_path)


@pytest.mark.asyncio
async def test_write_docx_creates_file():
    """write_docx creates a .docx and returns its path."""
    from app.tools.docx_tool import write_docx
    output = os.path.join(tempfile.gettempdir(), "test_write_output.docx")
    try:
        path_out = os.path.join(tempfile.gettempdir(), "test_write_output.docx")
        result = await write_docx({
            "file_path": path_out,
            "title": "Test Document",
            "paragraphs": [
                {"text": "Paragraph one", "style": "Normal"},
                {"text": "Paragraph two", "style": "Normal"},
            ],
        }, {})

        assert result["error"] is None
        assert result["type"] == "document"
        assert result["tool"] == "write_docx"

        # Verify file was created
        written_path = result["metadata"].get("file_path", output)
        assert os.path.exists(written_path)
        assert os.path.getsize(written_path) > 0
    finally:
        for p in [output, result.get("metadata", {}).get("file_path", "")]:
            if p and os.path.exists(p):
                os.unlink(p)


def _create_minimal_docx() -> str:
    """Create a minimal .docx file for testing. Returns path or None."""
    try:
        from docx import Document
        path = os.path.join(tempfile.gettempdir(), "test_project_agent.docx")
        doc = Document()
        doc.add_paragraph("Hello, Project Agent docx test!")
        doc.add_heading("Test Section", level=1)
        doc.add_paragraph("This is a test paragraph with some content.")
        doc.save(path)
        return path
    except ImportError:
        return None
