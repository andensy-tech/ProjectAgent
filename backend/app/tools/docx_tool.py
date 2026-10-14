"""
Word (.docx) tool — read and write Word documents using python-docx.
"""

import os
import logging
from typing import Dict, Any

logger = logging.getLogger(__name__)


async def read_docx(params: Dict[str, Any], context: Dict[str, Any]) -> Dict[str, Any]:
    """Extract text and tables from a .docx file.

    Args:
        params: {"file_path": str, "extract_tables": bool (default True)}
        context: Execution context

    Returns:
        {"type": "document", "tool": "read_docx", "content": str,
         "metadata": {file_name, file_size, paragraph_count, table_count}}
    """
    from docx import Document

    file_path = params.get("file_path", "")
    if not file_path or not os.path.exists(file_path):
        return {"type": "document", "tool": "read_docx", "content": "", "metadata": {}, "error": f"File not found: {file_path}"}

    try:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        content = "\n".join(paragraphs)

        metadata = {
            "file_name": os.path.basename(file_path),
            "file_size": os.path.getsize(file_path),
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
        }

        if params.get("extract_tables", True) and doc.tables:
            table_data = []
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                table_data.append(rows)
            metadata["tables"] = table_data

        return {
            "type": "document",
            "tool": "read_docx",
            "content": content.strip() or "(empty document)",
            "metadata": metadata,
            "error": None,
        }
    except Exception as e:
        logger.error("read_docx failed: %s", e)
        return {"type": "document", "tool": "read_docx", "content": "", "metadata": {}, "error": str(e)}


async def write_docx(params: Dict[str, Any], context: Dict[str, Any]) -> Dict[str, Any]:
    """Create a .docx file from structured data.

    Args:
        params: {"file_path": str, "title": str, "paragraphs": [{"text": str, "style": str, "level": int}],
                 "tables": [[[cell_value]]]}
        context: Execution context

    Returns:
        {"type": "document", "tool": "write_docx", "content": "file_path",
         "metadata": {file_name, file_size, paragraph_count, table_count}}
    """
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    file_path = params.get("file_path", "")
    if not file_path:
        return {"type": "document", "tool": "write_docx", "content": "", "metadata": {}, "error": "file_path is required"}

    try:
        doc = Document()

        # Title
        title = params.get("title", "")
        if title:
            heading = doc.add_heading(title, level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Paragraphs
        para_count = 0
        for p in params.get("paragraphs", []):
            text = p.get("text", "")
            style = p.get("style", "Normal")
            level = p.get("level")

            if level is not None:
                doc.add_heading(text, level=level)
            elif style == "Heading 1":
                doc.add_heading(text, level=1)
            elif style == "Heading 2":
                doc.add_heading(text, level=2)
            elif style == "Heading 3":
                doc.add_heading(text, level=3)
            elif style == "List Bullet":
                doc.add_paragraph(text, style="List Bullet")
            elif style == "List Number":
                doc.add_paragraph(text, style="List Number")
            else:
                doc.add_paragraph(text)
            para_count += 1

        # Tables
        table_count = 0
        for table_def in params.get("tables", []):
            if not table_def:
                continue
            rows_count = len(table_def)
            cols_count = max(len(r) for r in table_def) if table_def else 0
            if cols_count == 0:
                continue
            table = doc.add_table(rows=rows_count, cols=cols_count, style="Light Grid Accent 1")
            for i, row_def in enumerate(table_def):
                for j, cell_val in enumerate(row_def):
                    table.cell(i, j).text = str(cell_val)
            table_count += 1
            doc.add_paragraph()  # spacing after table

        os.makedirs(os.path.dirname(file_path) or ".", exist_ok=True)
        doc.save(file_path)

        return {
            "type": "document",
            "tool": "write_docx",
            "content": file_path,
            "metadata": {
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path) if os.path.exists(file_path) else 0,
                "paragraph_count": para_count,
                "table_count": table_count,
            },
            "error": None,
        }
    except Exception as e:
        logger.error("write_docx failed: %s", e)
        return {"type": "document", "tool": "write_docx", "content": "", "metadata": {}, "error": str(e)}
