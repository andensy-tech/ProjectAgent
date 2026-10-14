"""
Excel (.xlsx) tool — read and write spreadsheets using openpyxl.
"""

import os
import logging
from typing import Dict, Any, List

logger = logging.getLogger(__name__)


async def read_xlsx(params: Dict[str, Any], context: Dict[str, Any]) -> Dict[str, Any]:
    """Read an Excel file and return its content as structured data.

    Args:
        params: {"file_path": str, "sheet_name": str (optional, default first sheet),
                 "header_row": int (0-based, default 0), "max_rows": int (optional)}
        context: Execution context

    Returns:
        {"type": "document", "tool": "read_xlsx", "content": str (summary),
         "metadata": {file_name, file_size, sheet_names, active_sheet, rows, columns, data: [[cells]]}}
    """
    import openpyxl

    file_path = params.get("file_path", "")
    if not file_path or not os.path.exists(file_path):
        return {"type": "document", "tool": "read_xlsx", "content": "", "metadata": {}, "error": f"File not found: {file_path}"}

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet_name = params.get("sheet_name") or wb.sheetnames[0]
        if sheet_name not in wb.sheetnames:
            wb.close()
            return {"type": "document", "tool": "read_xlsx", "content": "", "metadata": {}, "error": f"Sheet not found: {sheet_name}"}

        ws = wb[sheet_name]
        header_row = params.get("header_row", 0)
        max_rows = params.get("max_rows")

        # Cache sheet names before potential close
        all_sheet_names = list(wb.sheetnames)

        rows_data: List[List[str]] = []
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if max_rows and row_count >= max_rows:
                break
            rows_data.append([str(c) if c is not None else "" for c in row])
            row_count += 1

        wb.close()

        headers = rows_data[header_row] if rows_data and header_row < len(rows_data) else []
        data_rows = rows_data[header_row + 1:] if header_row < len(rows_data) else []

        # Build summary content
        summary_parts = [
            f"Sheet: {sheet_name}",
            f"Rows: {len(data_rows)}, Columns: {len(headers) if headers else 'N/A'}",
        ]
        if headers:
            summary_parts.append(f"Headers: {', '.join(headers[:20])}{'...' if len(headers) > 20 else ''}")
        if data_rows:
            summary_parts.append(f"\nFirst {min(5, len(data_rows))} rows:")
            for i, row in enumerate(data_rows[:5]):
                preview = ", ".join(str(c) for c in row[:6])
                summary_parts.append(f"  Row {i + 1}: {preview}{'...' if len(row) > 6 else ''}")

        return {
            "type": "document",
            "tool": "read_xlsx",
            "content": "\n".join(summary_parts),
            "metadata": {
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "sheet_names": all_sheet_names,
                "active_sheet": sheet_name,
                "rows": len(data_rows),
                "columns": len(headers),
                "headers": headers,
                "data": data_rows[:100],  # Cap at 100 rows
            },
            "error": None,
        }
    except Exception as e:
        logger.error("read_xlsx failed: %s", e)
        return {"type": "document", "tool": "read_xlsx", "content": "", "metadata": {}, "error": str(e)}


async def write_xlsx(params: Dict[str, Any], context: Dict[str, Any]) -> Dict[str, Any]:
    """Create an Excel file from structured data.

    Args:
        params: {"file_path": str, "sheet_name": str (default "Sheet1"),
                 "headers": [str], "rows": [[cell_value]]}
        context: Execution context

    Returns:
        {"type": "document", "tool": "write_xlsx", "content": "file_path",
         "metadata": {file_name, file_size, sheet_name, rows, columns}}
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    file_path = params.get("file_path", "")
    if not file_path:
        return {"type": "document", "tool": "write_xlsx", "content": "", "metadata": {}, "error": "file_path is required"}

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = params.get("sheet_name", "Sheet1")

        # Write headers
        headers = params.get("headers", [])
        if headers:
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=str(h))
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.font = Font(bold=True, color="FFFFFF")

        # Write data rows
        rows = params.get("rows", [])
        for row_idx, row_data in enumerate(rows, start=2):
            for col_idx, cell_value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=cell_value)

        # Auto-adjust column widths (basic)
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 50)

        os.makedirs(os.path.dirname(file_path) or ".", exist_ok=True)
        wb.save(file_path)

        return {
            "type": "document",
            "tool": "write_xlsx",
            "content": file_path,
            "metadata": {
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path) if os.path.exists(file_path) else 0,
                "sheet_name": ws.title,
                "rows": len(rows),
                "columns": max(len(headers), max((len(r) for r in rows), default=0)),
            },
            "error": None,
        }
    except Exception as e:
        logger.error("write_xlsx failed: %s", e)
        return {"type": "document", "tool": "write_xlsx", "content": "", "metadata": {}, "error": str(e)}
